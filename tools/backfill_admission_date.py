"""입원일(actual_admission_date) 백필 — bokju-crm.

`header_index_map`이 '입원일 / 비고' 컬럼을 못 찾던 시절에 적재된 상담 행은
입원완료인데도 실제 입원일이 비어 있다. /ward 재원 명단은 입원일이 있는 건만
보여주므로, 그 환자들이 명부에서 통째로 빠진다.

이 도구는 **UPDATE 전용**이다. 환자도 상담도 새로 만들지 않는다 —
excel_import 를 --apply 로 다시 돌리면 중복이 섞여 들어오기 때문에
(도구가 완전 멱등이 아니다) 이미 있는 행의 빈 칸만 메우는 쪽을 택했다.

사용법:
    python tools/backfill_admission_date.py "/data/상담내역 종합.xlsx"
        → dry-run. 무엇을 채울지만 출력하고 DB는 건드리지 않는다.
    python tools/backfill_admission_date.py ... --apply
        → 실제 UPDATE. 직전 자동 백업.
    --overwrite 를 주면 이미 값이 있는 칸도 엑셀 값으로 덮는다 (기본은 빈 칸만).
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

import models  # noqa: E402
from tools.excel_import import (  # noqa: E402
    detect_schema, header_index_map, parse_date, norm_str,
    normalize_admission_status, backup_db, EXCLUDED_SHEETS,
)

# 스키마별 입원일 컬럼 이름. A는 '예정일'(입원완료면 실제 입원일로 본다),
# B/C는 병합 그룹 라벨이라 header_index_map(label_rows=...) 이 있어야 잡힌다.
ADMIT_DATE_HEADERS = ("입원일 / 비고", "입원일/비고", "예정일")


def sheet_rows(ws):
    """시트 → (스키마, 입원완료 행의 (이름, 상담일, 입원일) 목록)."""
    top = list(ws.iter_rows(min_row=1, max_row=5, values_only=True))
    schema, hi = detect_schema(top)
    if schema is None:
        return None, []
    header_row = top[hi - 1]
    label_rows = [top[i] for i in (hi - 2, hi) if 0 <= i < len(top)]
    idx = header_index_map(header_row, label_rows)

    date_col = next((idx[h] for h in ADMIT_DATE_HEADERS if h in idx), None)
    name_col, st_col, cd_col = (idx.get("환자이름"), idx.get("입원여부"),
                                idx.get("상담일자"))
    if date_col is None or None in (name_col, st_col, cd_col):
        return schema, []

    out = []
    for r in ws.iter_rows(min_row=hi + 1, values_only=True):
        if max(name_col, st_col, cd_col, date_col) >= len(r):
            continue
        if normalize_admission_status(r[st_col]) != "입원완료":
            continue
        name = norm_str(r[name_col])
        consult_date = parse_date(r[cd_col])
        admit_date = parse_date(r[date_col])
        if not (name and consult_date and admit_date):
            continue
        out.append((name, consult_date, admit_date))
    return schema, out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx_path", help="상담내역 종합 엑셀 경로")
    ap.add_argument("--apply", action="store_true", help="실제 DB에 UPDATE")
    ap.add_argument("--overwrite", action="store_true",
                    help="이미 입원일이 있는 행도 엑셀 값으로 덮어쓴다")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx_path, read_only=True, data_only=True)
    conn = models.get_db()
    stat = Counter()
    updates, ambiguous, unmatched, conflicts = [], [], [], []

    for name in wb.sheetnames:
        if name in EXCLUDED_SHEETS:
            continue
        schema, rows = sheet_rows(wb[name])
        if schema is None:
            print(f"  [스키마 감지 실패] {name}")
            continue
        for pname, consult_date, admit_date in rows:
            stat["엑셀 입원완료(날짜 있음)"] += 1
            # excel_import 의 중복 판정과 같은 키: 환자 + 상담일
            hits = conn.execute(
                """SELECT c.id, c.actual_admission_date
                     FROM consultations c JOIN patients p ON p.id = c.patient_id
                    WHERE p.name = ? AND c.consult_date = ?""",
                (pname, consult_date),
            ).fetchall()
            if not hits:
                stat["DB에 없음"] += 1
                unmatched.append((name, pname, consult_date))
                continue
            if len(hits) > 1:
                # 동명이인이 같은 날 상담 — 어느 행인지 특정 불가. 건드리지 않는다.
                stat["동명이인 모호"] += 1
                ambiguous.append((name, pname, consult_date, len(hits)))
                continue
            cid, cur = hits[0][0], (hits[0][1] or "").strip()
            if cur and cur == admit_date:
                stat["이미 동일"] += 1
            elif cur and not args.overwrite:
                stat["값 다름(보존)"] += 1
                conflicts.append((pname, consult_date, cur, admit_date))
            else:
                stat["채울 대상"] += 1
                updates.append((admit_date, cid))

    print("=" * 60)
    for k, v in stat.items():
        print(f"  {k:22s} {v:5d}")
    print("=" * 60)

    if conflicts:
        print(f"\n[DB값 != 엑셀값 {len(conflicts)}건 — 기본은 DB 보존, --overwrite 로 덮음]")
        for c in conflicts[:10]:
            print(f"  {c[0]} ({c[1]}): DB={c[2]}  엑셀={c[3]}")
    if ambiguous:
        print(f"\n[동명이인·같은날 {len(ambiguous)}건 — 수동 확인 필요]")
        for a in ambiguous[:10]:
            print(f"  {a[1]} {a[2]} ({a[0]}) 후보 {a[3]}건")
    if unmatched:
        print(f"\n[DB에 없는 행 {len(unmatched)}건 — 적재 때 스킵된 행]")
        for u in unmatched[:10]:
            print(f"  {u[1]} {u[2]} ({u[0]})")

    if not args.apply:
        print(f"\n※ dry-run. {len(updates)}건을 채울 예정. 실제 반영은 --apply")
        conn.close()
        return

    conn.close()
    backup_db()
    conn = models.get_db()
    try:
        conn.executemany(
            "UPDATE consultations SET actual_admission_date=? WHERE id=?", updates)
        conn.commit()
    finally:
        conn.close()
    print(f"\n완료: {len(updates)}건 입원일 반영")


if __name__ == "__main__":
    main()
