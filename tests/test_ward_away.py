import os
import tempfile
import unittest
from datetime import date, timedelta
from unittest.mock import patch

import app as main
import models


class AwayManagementTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        db = patch.object(models, 'DB_PATH', os.path.join(self.tmp.name, 'test.db'))
        db.start(); self.addCleanup(db.stop)
        models.init_db()
        models.ensure_admin_user('away-test', 'test-password', display_name='테스트')
        self.uid = models.get_user('away-test')['id']
        boot = patch.object(main, '_db_initialized', True)
        boot.start(); self.addCleanup(boot.stop)
        main.app.config.update(TESTING=True)
        self.client = main.app.test_client()
        self.login(2)
        with models.get_db() as conn:
            for pid in (1, 2):
                conn.execute('INSERT INTO patients (id, name, gender) VALUES (?, ?, ?)',
                             (pid, f'테스트환자{pid}', 'F'))
                conn.execute('''INSERT INTO consultations
                    (id, patient_id, consult_date, admission_status, actual_admission_date,
                     patient_age, primary_diagnosis, discharge_due_date)
                    VALUES (?, ?, '2026-01-01', '입원완료', '2026-01-01', 70, '테스트병명', '2026-12-31')''',
                             (pid, pid))
            for cid, kind, day, returned in [
                (1, '응급전원', '2026-02-01', '2026-02-05'),
                (1, '응급전원', '2026-03-01', None),
                (2, '응급전원', '2026-02-03', None),
                (2, '모병원 외래치료', '2026-01-10', '2026-01-10'),
                (2, '기타', '2026-01-11', None),
            ]:
                conn.execute('''INSERT INTO admission_events
                    (consultation_id, event_type, event_date, returned_at, hospital, memo)
                    VALUES (?, ?, ?, ?, '테스트기관', '<script>사유</script>')''',
                             (cid, kind, day, returned))

    def login(self, level):
        with self.client.session_transaction() as session:
            session.clear()
            session.update(user_id=self.uid, username='away-test', display_name='테스트',
                           role='staff', cooperation_permissions_v2=True,
                           perms={k: (level if k == 'ward' else 1) for k in main.MENU_KEYS})

    def test_unique_patients_repeat_transfers_and_discharged_history(self):
        with models.get_db() as conn:
            conn.execute("UPDATE consultations SET discharge_date='2026-03-10' WHERE id=2")
        rows = models.list_away_records(event_type='응급전원')
        stats = models.away_record_stats(rows)
        self.assertEqual((stats['patients'], stats['returned_patients'], stats['events']), (2, 1, 3))
        self.assertEqual(stats['open_patients'], 2)
        self.assertEqual(stats['patient_rate'], 50)
        self.assertEqual(stats['event_rate'], 33.3)
        self.assertEqual(len(models.list_away_records(date_from='2026-02-01', date_to='2026-02-28')), 2)
        self.assertEqual(models.away_record_stats([])['patient_rate'], 0)

    def test_tab_render_filters_and_monthly_cohort(self):
        page = self.client.get('/ward?tab=away')
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        for label in ('외진 환자', '퇴원일(전원일)', '테스트병명', '복귀 예정 저장', '70', '&lt;script&gt;'):
            self.assertIn(label, html)
        self.assertNotIn('<script>사유</script>', html)
        with main.app.test_request_context('/ward?tab=away&away_status=returned&away_q=환자1'):
            report = main._ward_away_report()
            self.assertEqual(len(report['rows']), 1)
            self.assertEqual(report['stats']['events'], 4)
            self.assertEqual(report['transfers']['returned_patients'], 1)
            self.assertEqual(dict(report['monthly'])['2026-02']['patients'], 2)
        for query in ('away_from=bad', 'away_from=2026-03-01&away_to=2026-02-01', 'away_type=invalid'):
            self.assertEqual(self.client.get('/ward?tab=away&' + query).status_code, 400)

    def test_list_and_feed_views_number_rows(self):
        """목록형이 기본이고, 피드형은 날짜별로 묶는다. 연번은 두 보기에서 같다."""
        html = self.client.get('/ward?tab=away').get_data(as_text=True)
        self.assertIn('class="on">☰ 목록형', html)
        self.assertIn('<td class="away-col-no" data-label="#">1</td>', html)
        self.assertIn('일째', html)          # 미복귀 기록의 '며칠째' 배지
        self.assertIn('번째 외진', html)     # 차수 표시
        self.assertIn('data-detail-active="false"', html)
        self.assertIn('/ward/away.xlsx', html)
        self.assertIn('미복귀 · 처리', html)
        self.assertIn('title="1번째 외진"', html)
        feed = self.client.get('/ward?tab=away&away_view=feed').get_data(as_text=True)
        self.assertIn('class="on">▤ 피드형', feed)
        self.assertIn('away-feed-date', feed)
        self.assertIn('<div class="away-card-no">1</div>', feed)
        self.assertNotIn('away-table', feed.split('<style>')[0])
        # 상세필터에 값이 있으면 펼친 채로 연다
        html = self.client.get('/ward?tab=away&away_gender=F').get_data(as_text=True)
        self.assertIn('data-detail-active="true"', html)
        with main.app.test_request_context('/ward?tab=away&away_view=bogus'):
            self.assertEqual(main._ward_away_report()['view'], 'list')

    def test_xlsx_export_follows_filters_and_permissions(self):
        from openpyxl import load_workbook
        import io
        resp = self.client.get('/ward/away.xlsx?away_status=returned&away_q=환자1')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp.mimetype)
        ws = load_workbook(io.BytesIO(resp.data)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ('연번', '차수', '환자'))
        self.assertEqual(rows[1][0], 1)
        self.assertEqual(rows[1][18], '복귀 완료')
        self.assertEqual(sum(1 for r in rows[1:] if r[0] is not None and isinstance(r[0], int)), 1)
        self.login(0)
        self.assertEqual(self.client.get('/ward/away.xlsx').status_code, 403)

    def test_transfer_outcome_is_not_a_return(self):
        """외진이 타 병원 전원으로 끝나면 복귀 통계에서 빠지고 단계는 퇴원이 된다."""
        self.login(2)
        url = '/api/admission-event/2/return'
        self.assertEqual(self.client.post(url, json={'return_outcome': '전원'}).status_code, 400)   # 병원 필수
        self.assertEqual(self.client.post(url, json={'return_outcome': '기타'}).status_code, 400)
        self.assertEqual(self.client.post(url, json={'return_outcome': '전원', 'return_hospital': '안동병원',
                                                     'return_room': '999호', 'return_date': '2026-03-05'}).status_code, 200)
        ev = models.get_admission_event(2)
        self.assertEqual((ev['returned_at'], ev['return_outcome'], ev['return_hospital'], ev['return_room']),
                         ('2026-03-05', '전원', '안동병원', None))
        pid = models.get_consultation(ev['consultation_id'])['patient_id']
        self.assertEqual(models.get_patient(pid)['lifecycle_stage'], '퇴원')
        stats = models.away_record_stats(models.list_away_records(event_type='응급전원'))
        self.assertEqual((stats['returned_events'], stats['transferred_events'], stats['open_events']), (1, 1, 1))
        with main.app.test_request_context('/ward?tab=away&away_status=transferred'):
            rows = main._ward_away_report()['rows']
            self.assertEqual([r['away_id'] for r in rows], [2])
        with main.app.test_request_context('/ward?tab=away&away_status=returned'):
            self.assertNotIn(2, [r['away_id'] for r in main._ward_away_report()['rows']])
        html = self.client.get('/ward?tab=away').get_data(as_text=True)
        self.assertIn('타 병원 전원</strong>', html)
        self.assertIn('안동병원', html)
        self.assertIn('타 병원 전원 1명 1건', html)

    def test_admission_date_falls_back_to_roster_episode(self):
        """상담에 입원일이 없어도 회차(원무 명부)의 입원일로 재원일수를 센다."""
        with models.get_db() as conn:
            conn.execute("UPDATE consultations SET actual_admission_date=NULL, admission_date=NULL WHERE id=1")
            conn.execute("""INSERT INTO admission_episodes (patient_id, consultation_id, episode_no, status, admitted_at, roster_key)
                            VALUES (1, 1, 1, 'admitted', '2026-01-10', 'r1')""")
            eid = conn.execute("SELECT id FROM admission_episodes WHERE consultation_id=1").fetchone()[0]
            conn.execute("UPDATE admission_events SET episode_id=? WHERE consultation_id=1", (eid,))
        with main.app.test_request_context('/ward?tab=away&away_q=환자1'):
            rows = main._ward_away_report()['rows']
        self.assertTrue(rows)
        self.assertTrue(all(r['admitted_on'] == '2026-01-10' for r in rows))
        self.assertTrue(all(r['stay_days_inclusive'] is not None for r in rows))

    def test_insights_group_by_reason_dx_hospital_and_days(self):
        with models.get_db() as conn:
            conn.execute("UPDATE admission_events SET memo='폐렴 | 명부 메모' WHERE id=1")
        rows = models.list_away_records()
        ins = main._away_insights(rows)
        self.assertEqual(ins['events'], len(rows))
        reasons = {b['label']: b for b in ins['by_reason']}
        self.assertIn('폐렴', reasons)                       # '|' 뒤 출처 메모는 떼고 묶는다
        self.assertIn('테스트기관', {b['label'] for b in ins['by_hospital']})
        self.assertTrue(all(b['events'] >= 1 for b in ins['by_type']))
        self.assertEqual(sum(d['count'] for d in ins['distribution']), ins['returned_n'])
        if ins['returned_n']:
            self.assertIsNotNone(ins['avg_days'])
        self.assertIsNone(main._away_insights([]))
        html = self.client.get('/ward?tab=away').get_data(as_text=True)
        self.assertIn('외진 환자 특성', html)
        self.assertIn('평균 복귀 소요', html)

    def test_return_persistence_validation_and_permissions(self):
        url = '/api/admission-event/2/return'
        self.login(1)
        html = self.client.get('/ward?tab=away').get_data(as_text=True)
        self.assertNotIn('id="away-register"', html)
        self.assertNotIn('복귀 예정 저장', html)
        self.assertEqual(self.client.post(url, json={}).status_code, 403)
        self.login(0)
        self.assertEqual(self.client.get('/ward?tab=away').status_code, 403)
        self.login(2)
        for day in ('2026-02-01', (date.today() + timedelta(days=1)).isoformat(), 'bad'):
            self.assertEqual(self.client.post(url, json={'return_date': day}).status_code, 400)
        self.assertEqual(self.client.post('/api/admission-event/5/return', json={}).status_code, 400)
        self.assertEqual(self.client.post(url, json={'return_date': '2026-03-05'}).status_code, 200)
        self.assertEqual(models.get_admission_event(2)['returned_at'], '2026-03-05')
        stats = models.away_record_stats(models.list_away_records(event_type='응급전원'))
        self.assertEqual((stats['returned_patients'], stats['returned_events']), (1, 2))
        self.assertEqual(self.client.post(url, json={}).status_code, 400)
        result = self.client.post('/api/consult/1/admission-event', json={
            'event_type': '응급전원', 'event_date': '2026-03-10', 'hospital': '병원', 'memo': '검사'})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(self.client.post('/api/consult/1/admission-event', json={
            'event_type': '응급전원', 'event_date': '2026-03-11'}).status_code, 400)

    def test_return_saves_room_note_and_moves_current_room(self):
        models.sync_admission_episode(1)
        response = self.client.post('/api/admission-event/2/return', json={
            'return_date': '2026-03-05', 'return_room': '301호', 'return_note': '산소 유지 필요'})
        self.assertEqual(response.status_code, 200)
        event = models.get_admission_event(2)
        self.assertEqual((event['return_room'], event['return_note']), ('301호', '산소 유지 필요'))
        # 기록만이 아니라 현재 병실도 따라가야 병상 화면이 어긋나지 않는다
        self.assertEqual(models.get_consultation(1)['room_number'], '301호')
        with models.get_db() as conn:
            episode = conn.execute(
                'SELECT room_number FROM admission_episodes WHERE consultation_id = 1').fetchone()
        self.assertEqual(episode['room_number'], '301호')
        row = [r for r in models.list_away_records() if r['away_id'] == 2][0]
        self.assertEqual(row['return_room'], '301호')

    def test_return_without_room_keeps_current_room(self):
        with models.get_db() as conn:
            conn.execute("UPDATE consultations SET room_number='205호' WHERE id=1")
        self.assertEqual(self.client.post('/api/admission-event/2/return',
                                          json={'return_date': '2026-03-05'}).status_code, 200)
        self.assertEqual(models.get_consultation(1)['room_number'], '205호')

    def test_return_rejects_overlong_room_and_note(self):
        for payload in ({'return_room': '방' * 51}, {'return_note': '가' * 3001}):
            self.assertEqual(self.client.post('/api/admission-event/2/return',
                                              json=payload).status_code, 400)
        self.assertIsNone(models.get_admission_event(2)['returned_at'])

    def test_details_saves_birth_date_and_readmission(self):
        self.login(1)
        self.assertEqual(self.client.post('/api/admission-event/2/details',
                                          json={'readmission': '예'}).status_code, 403)
        self.login(2)
        response = self.client.post('/api/admission-event/2/details',
                                    json={'birth_date': '1955-04-02', 'readmission': '예'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(models.get_admission_event(2)['readmission'], '예')
        # 생년월일은 환자에 저장되므로 같은 환자의 다른 외진 기록에도 보인다
        rows = {r['away_id']: r for r in models.list_away_records()}
        self.assertEqual(rows[2]['birth_date'], '1955-04-02')
        self.assertEqual(rows[1]['birth_date'], '1955-04-02')
        self.assertIsNone(rows[1]['readmission'])

    def test_details_rejects_bad_values_and_non_away_event(self):
        future = (date.today() + timedelta(days=1)).isoformat()
        for payload in ({'birth_date': future}, {'birth_date': '55-4-2'}, {'readmission': '아마도'}):
            self.assertEqual(self.client.post('/api/admission-event/2/details',
                                              json=payload).status_code, 400)
        # 5번은 '기타' 이벤트 — 외진 명부 항목이 아니다
        self.assertEqual(self.client.post('/api/admission-event/5/details',
                                          json={'readmission': '예'}).status_code, 404)
        self.assertIsNone(models.get_patient(1)['birth_date'])

    def test_lifetime_numbering_survives_filters_and_readmission(self):
        with models.get_db() as conn:
            conn.execute("INSERT INTO consultations (id, patient_id, consult_date) VALUES (3, 1, '2026-04-01')")
            conn.execute("INSERT INTO admission_events (consultation_id,event_type,event_date) VALUES (3,'모병원 외래치료','2026-04-03')")
            conn.execute("INSERT INTO admission_events (consultation_id,event_type) VALUES (3,'응급전원')")
        rows = models.list_away_records(date_from='2026-03-01', event_type='응급전원')
        self.assertEqual([(r['away_id'], r['away_number']) for r in rows], [(2, 2)])
        rows = models.list_away_records(date_from='2026-04-01')
        self.assertEqual(rows[0]['away_number'], 3)
        missing = [r for r in models.list_away_records() if not r['event_date']]
        self.assertIsNone(missing[0]['away_number'])

    def test_inclusive_days_and_roster_filters(self):
        with main.app.test_request_context('/ward?tab=away&away_number=2&away_age_min=70&away_age_max=70&away_gender=F&away_hospital=기관&away_dx=병명&away_days_min=2'):
            report = main._ward_away_report()
        self.assertEqual({r['away_id'] for r in report['rows']}, {2, 3})
        self.assertEqual(report['stats']['events'], 4)
        row = next(r for r in report['rows'] if r['away_id'] == 2)
        self.assertEqual(row['stay_days_inclusive'], (date.today() - date(2026, 1, 1)).days + 1)
        self.assertEqual(row['away_days_inclusive'], (date.today() - date(2026, 3, 1)).days + 1)
        with models.get_db() as conn:
            conn.execute("UPDATE consultations SET discharge_date='2026-03-10' WHERE id=2")
        with main.app.test_request_context('/ward?tab=away'):
            rows = main._ward_away_report()['rows']
        returned = next(r for r in rows if r['away_id'] == 4)
        self.assertEqual(returned['away_days_inclusive'], 1)
        self.assertEqual(returned['stay_days_inclusive'], 69)
        for query in ('away_gender=M', 'away_age_min=71', 'away_age_max=69', 'away_hospital=없는기관',
                      'away_dx=없는병명', 'away_number=9', 'away_days_min=99999', 'away_dday_max=-99999'):
            with main.app.test_request_context('/ward?tab=away&' + query):
                self.assertEqual(main._ward_away_report()['rows'], [], query)
        for query in ('away_number=0', 'away_age_min=-1', 'away_age_min=80&away_age_max=70',
                      'away_days_min=oops', 'away_gender=bad', 'away_dday_max=1.5'):
            self.assertEqual(self.client.get('/ward?tab=away&' + query).status_code, 400, query)

    def _today_admission_counts(self):
        s = models.dashboard_summary()['summary']
        return s['admission_today_planned'], s['admission_today_completed'], s['admission_today']

    def test_expected_return_counts_as_planned_then_completed_admission(self):
        today = date.today().isoformat()
        # 이 픽스처의 재원 상담엔 입원예정일이 없어 대시보드 '오늘 입원'은 0에서 시작한다.
        self.assertEqual(self._today_admission_counts(), (0, 0, 0))
        # 미복귀 외진(#3, 상담 2)에 복귀 예정일=오늘 → 입원 예정 1
        r = self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': today})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self._today_admission_counts(), (1, 0, 1))
        data = models.dashboard_summary()
        row = next(x for x in data['admission_by_status']['planned'] if x.get('admission_kind') == 'return')
        self.assertEqual((row['id'], row['away_event_id'], row['admission_bucket_label']), (2, 3, '복귀 예정'))
        self.assertEqual(models.away_now([2])[0]['expected_return_date'], today)
        # 재원 관리 [복귀 예정]은 병실·기타 사항을 미리 받아 두기만 한다 — 아직 병실 이동 없음
        r = self.client.post('/api/admission-event/3/expected-return',
                             json={'expected_return_date': today, 'return_room': '702호', 'return_note': '휠체어'})
        self.assertEqual(r.status_code, 200, r.get_json())
        ev = models.get_admission_event(3)
        self.assertEqual((ev['return_room'], ev['return_note'], ev['returned_at']), ('702호', '휠체어', None))
        self.assertIsNone(models.get_consultation(2)['room_number'])
        # 대시보드 입원 처리 [완료] = 날짜 없이 /return → 오늘 복귀, 예정에서 빠지고 완료 1, 병실 이동
        r = self.client.post('/api/admission-event/3/return', json={})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(r.get_json()['return_date'], today)
        self.assertEqual(self._today_admission_counts(), (0, 1, 1))
        self.assertEqual(models.get_consultation(2)['room_number'], '702호')
        self.assertEqual(models.get_admission_event(3)['return_note'], '휠체어')
        # 복귀 처리된 뒤엔 예정일을 못 바꾼다
        r = self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': today})
        self.assertEqual(r.status_code, 400)
        # 복귀 취소 → 다시 외진 중 + 복귀 예정(오늘) = 입원 예정 1, 완료 0. 병실·기타 사항은 유지
        r = self.client.post('/api/admission-event/3/return/undo', json={'expected_return_date': today})
        self.assertEqual(r.status_code, 200, r.get_json())
        ev = models.get_admission_event(3)
        self.assertEqual((ev['returned_at'], ev['return_outcome'], ev['expected_return_date'], ev['return_room']),
                         (None, None, today, '702호'))
        self.assertEqual(self._today_admission_counts(), (1, 0, 1))
        self.assertEqual(self.client.post('/api/admission-event/3/return/undo', json={}).status_code, 400)  # 미복귀
        # 다시 완료
        self.assertEqual(self.client.post('/api/admission-event/3/return', json={}).status_code, 200)
        self.assertEqual(self._today_admission_counts(), (0, 1, 1))
        # 타 병원 전원으로 종결된 외진(#2, 상담 1)은 '완료'가 아니다
        r = self.client.post('/api/admission-event/2/return',
                             json={'return_date': today, 'return_outcome': '전원', 'return_hospital': '타병원'})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self._today_admission_counts(), (0, 1, 1))
        self.assertEqual(self.client.post('/api/admission-event/2/return/undo', json={}).status_code, 400)  # 전원은 불가

    def test_expected_return_validation_and_create(self):
        # 외진일보다 빠른 복귀 예정일 거부
        r = self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': '2026-02-01'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('빠릅니다', r.get_json()['error'])
        r = self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': '2026/02/09'})
        self.assertEqual(r.status_code, 400)
        # 외진이 아닌 '기타' 이벤트(#5)엔 복귀 예정일이 없다
        r = self.client.post('/api/admission-event/5/expected-return', json={'expected_return_date': '2026-02-09'})
        self.assertEqual(r.status_code, 404)
        # 빈 값은 '미정'으로 되돌린다
        self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': '2026-02-09'})
        r = self.client.post('/api/admission-event/3/expected-return', json={'expected_return_date': ''})
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(models.get_admission_event(3)['expected_return_date'])
        # 새 외진 기록에 복귀 예정일을 함께 저장 (상담 3, 미복귀 없음)
        with models.get_db() as conn:
            conn.execute("INSERT INTO consultations (id, patient_id, consult_date, admission_status, actual_admission_date) "
                         "VALUES (3, 1, '2026-04-01', '입원완료', '2026-04-01')")
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        r = self.client.post('/api/consult/3/admission-event', json={
            'event_type': '모병원 외래치료', 'event_date': date.today().isoformat(),
            'hospital': '모병원', 'memo': '외래', 'expected_return_date': tomorrow})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(models.get_admission_event(r.get_json()['id'])['expected_return_date'], tomorrow)
        self.assertEqual(self._today_admission_counts(), (0, 0, 0))   # 내일 예정이라 오늘엔 안 잡힌다
        self.assertEqual(models.dashboard_summary()['summary']['admission_planned_week'], 1)


if __name__ == '__main__':
    unittest.main()
