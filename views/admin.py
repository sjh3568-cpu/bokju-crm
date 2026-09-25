"""사용자 관리(어드민) — 계정·권한 매트릭스·감사 로그·권한 요청.

app.py에서 분리(2026-09-15). 라우트 경로·동작은 그대로, 엔드포인트 이름만 "admin.<함수>"가 됐다.
"""
import csv
import hashlib
import hmac
import time
import io
import json
import logging
import os
import re
import secrets
import threading
import calendar
from contextlib import closing
from functools import lru_cache
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlencode, urlsplit
from dotenv import load_dotenv
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from werkzeug.security import check_password_hash
from flask import (
    Flask, abort, current_app, flash, g, jsonify, redirect, render_template,
    request, send_file, session, url_for,
)
import backup
import dashboard_metrics
import models
import release_notes
from auth import (
    admin_required, authenticate, current_user, is_locked_out,
    login_required, login_user, logout_user, menu_level,
)
from config import (
    APP_VERSION, APP_DEVELOPER, CALL_RECORDING_URL,
    ACTIVITY_ACTIVE_OPTIONS, ACTIVITY_DIAPER_OPTIONS, ACTIVITY_OTHERS_OPTIONS,
    ACTIVITY_WHEELCHAIR_OPTIONS, ADMISSION_DOCS, ADMISSION_STATUSES,
    AUDIT_ACTION_LABELS, AUDIT_CATEGORIES, AUDIT_CATEGORY_OTHER,
    AUDIT_CRITICAL_ACTIONS, AUDIT_RETENTION_DAYS,
    ADMISSION_EVENT_TYPES, ATTENDING_DOCTORS, DOCTOR_DEPT_CODES,
    BED_OPTIONS, CAREGIVER_OPTIONS,
    CONSCIOUSNESS_MAIN_OPTIONS, CONSULT_CHANNELS, CONVERSATION_LEVEL_OPTIONS,
    CONSULT_RESULTS, CONSULT_RESULT_REASON_LABELS, REJECTION_REASONS,
    COMM_CHANNELS, COMM_INBOUND_CHANNELS,
    COST_GUIDANCE_OPTIONS, CURRENT_LOCATION_TYPES, DIET_TYPES, DIET_LAYOUT,
    DISEASES_CHECKLIST, DISEASES_GROUPS, GUARDIAN_RELATION_SUGGESTIONS,
    HEARING_OPTIONS, INFO_PROVIDED_OPTIONS,
    COUNSELORS, COUNSELORS_ACTIVE, DISEASES_LAYOUT, OTHERS_LAYOUT, ROOM_CAPACITY,
    ROOM_BED_CAPACITIES,
    WARDS, MGMT_TAG_PRESETS,
    ROLE_LABELS, SEED_USERS,
    MENUS, MENU_KEYS, MENU_MAX_LEVEL, ROLE_PRESETS, role_preset,
    PERM_HIDDEN, PERM_VIEW, PERM_EDIT, PERM_CREATE,
    PERM_LEVELS, PERM_LEVEL_LABELS,
    INSURANCE_TYPES, OTHERS_CHECKLIST, REFERRAL_SOURCE_GROUPS, REFERRAL_TYPES,
    INBOUND_CHANNEL_REFERRAL, INBOUND_CHANNEL_LABELS,
    STAFF_REFERRAL_ORGS, STAFF_REFERRAL_DEPTS,
    LIFECYCLE_STAGES, LIFECYCLE_EVENT_TYPES, LEGACY_STAGE_MAP, CARE_PHASES,
    SIDO_LIST, SIGUNGU_INDEX, SIGUNGU_LIST,
    SMS_TEMPLATE_GROUPS, SMS_PLACEHOLDERS,
    SPECIAL_CARE_OPTIONS, SPECIAL_CARE_NOTE_FIELDS,
    THERAPY_OPTIONS, TRANSPORT_OPTIONS, WOUND_CARE_OPTIONS, WOUND_CARE_NOTE_FIELDS,
)
import sms as sms_gateway
import partnerships
import support_requests
import transport
import homepage_board
import ward_moves
from flask import Blueprint
import logging
from app import (  # noqa: E402 — app.py 공용 헬퍼·상수 (app.py 맨 아래에서 이 모듈을 import하므로 순환 없음)
    _MIN_PW_LEN,
    app,
)

logger = logging.getLogger(__name__)
bp = Blueprint("admin", __name__)

# ───────────────────── 사용자관리(어드민전용) ─────────────────────

_VALID_ROLES = {"admin", "staff", "viewer"}

AUDIT_PAGE_SIZE = 100  # 이력 관리 페이지당 행 수

AUDIT_PAGE_SIZE_OPTIONS = (50, 100, 200, 500)

AUDIT_EXPORT_LIMIT = 20000  # CSV 한 번에 내보내는 최대 행 수

def _audit_filters_from_request():
    """이력 관리 화면의 필터를 요청에서 읽어 models 인자 형태로 정리."""
    category = (request.args.get("category") or "").strip()
    action = (request.args.get("action") or "").strip()
    known = [a for _, acts in AUDIT_CATEGORIES.values() for a in acts]
    if action:
        actions, exclude = [action], None
    elif category == AUDIT_CATEGORY_OTHER:
        actions, exclude = None, known          # 어느 분류에도 없는 action
    elif category in AUDIT_CATEGORIES:
        actions, exclude = AUDIT_CATEGORIES[category][1], None
    else:
        category, actions, exclude = "", None, None
    return {
        "date_from": (request.args.get("from") or "").strip() or None,
        "date_to": (request.args.get("to") or "").strip() or None,
        "username": (request.args.get("user") or "").strip() or None,
        "target_type": (request.args.get("target") or "").strip() or None,
        "q": (request.args.get("q") or "").strip() or None,
        "actions": actions,
        "exclude_actions": exclude,
    }, category, action

@bp.route("/admin/audit")
@admin_required
def admin_audit():
    """이력 관리 — 누가·언제·무엇을 조회/입력/수정/삭제했는지 (audit_log)."""
    filters, category, action = _audit_filters_from_request()
    try:
        page = max(1, int(request.args.get("page") or 1))
    except (ValueError, TypeError):
        page = 1
    try:
        requested = int(request.args.get("page_size") or AUDIT_PAGE_SIZE)
    except (ValueError, TypeError):
        requested = AUDIT_PAGE_SIZE
    page_size = requested if requested in AUDIT_PAGE_SIZE_OPTIONS else AUDIT_PAGE_SIZE

    total = models.count_audit_logs(**filters)
    total_pages = max(1, (total + page_size - 1) // page_size)
    page = min(page, total_pages)
    offset = (page - 1) * page_size
    rows = models.list_audit_logs(**filters, limit=page_size, offset=offset)
    # CSV 내보내기 링크 — 현재 검색 조건만 유지 (페이지·보기개수는 의미 없음)
    export_args = {k: v for k, v in request.args.items()
                   if k not in ("page", "page_size") and v}
    export_qs = ("?" + urlencode(export_args)) if export_args else ""
    return render_template(
        "audit.html", rows=rows, total=total, page=page, total_pages=total_pages,
        export_qs=export_qs,
        page_size=page_size, page_size_options=AUDIT_PAGE_SIZE_OPTIONS,
        page_start=offset,
        category=category, action=action,
        action_counts=models.audit_action_counts(**filters),
        users=models.audit_usernames(),
        target_types=models.audit_target_types(),
        span=models.audit_log_span(),
        AUDIT_ACTION_LABELS=AUDIT_ACTION_LABELS,
        AUDIT_CATEGORIES=AUDIT_CATEGORIES,
        AUDIT_CATEGORY_OTHER=AUDIT_CATEGORY_OTHER,
        AUDIT_CRITICAL_ACTIONS=AUDIT_CRITICAL_ACTIONS,
        AUDIT_RETENTION_DAYS=AUDIT_RETENTION_DAYS,
    )

@bp.route("/admin/audit/export")
@admin_required
def admin_audit_export():
    """현재 필터 조건의 이력을 CSV로 — 개인정보 열람기록 제출·내부 감사용."""
    filters, _category, _action = _audit_filters_from_request()
    rows = models.list_audit_logs(**filters, limit=AUDIT_EXPORT_LIMIT, offset=0)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["일시", "사용자ID", "아이디", "이름", "행위", "행위(코드)",
                "대상종류", "대상ID", "상세", "IP"])
    for r in rows:
        w.writerow([
            r.get("created_at") or "",
            r.get("user_id") if r.get("user_id") is not None else "",
            r.get("username") or "",
            r.get("display_name") or "",
            AUDIT_ACTION_LABELS.get(r.get("action"), r.get("action") or ""),
            r.get("action") or "",
            r.get("target_type") or "",
            r.get("target_id") if r.get("target_id") is not None else "",
            r.get("detail") or "",
            r.get("ip") or "",
        ])
    models.log_audit(
        user_id=g.user["id"], username=g.user["username"],
        action="export_csv", target_type="audit_log",
        detail=f"이력 {len(rows)}건", ip=request.remote_addr,
    )
    data = buf.getvalue().encode("utf-8-sig")  # Excel 한글 깨짐 방지
    return send_file(
        io.BytesIO(data), mimetype="text/csv", as_attachment=True,
        download_name=f"audit_log_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
    )

def _parse_perms_form(form):
    """폼의 perm_<menu> 값 → {menu: level} (메뉴별 지원 최대 레벨로 클램프)."""
    perms = {}
    for k in MENU_KEYS:
        try:
            lvl = int(form.get(f"perm_{k}", 0))
        except (TypeError, ValueError):
            lvl = 0
        perms[k] = max(0, min(MENU_MAX_LEVEL[k], lvl))
    return perms

def _count_user_managers(exclude_id=None):
    """사용자 관리(users) 권한이 '수정' 이상인 활성 계정 수. 마지막 관리자 보호용."""
    n = 0
    for u in models.list_users():
        if exclude_id is not None and u["id"] == exclude_id:
            continue
        if u["active"] and u["perms"].get("users", 0) >= PERM_EDIT:
            n += 1
    return n

@bp.route("/admin/users")
@admin_required
def admin_users():
    users = models.list_users()
    return render_template(
        "users.html", users=users, menus=MENUS, perm_labels=PERM_LEVEL_LABELS,
        role_presets=ROLE_PRESETS, menu_max=MENU_MAX_LEVEL,
        password_reset_requests=models.list_pending_password_reset_requests(),
        permission_requests=models.list_permission_requests(pending_only=True),
    )

@bp.route('/admin/permission-requests/<int:request_id>',methods=['POST'])
@admin_required
def admin_permission_request_resolve(request_id):
    status=(request.form.get('status') or '').strip()
    if status not in ('승인','반려'): abort(400)
    models.resolve_permission_request(request_id,g.user['id'],status)
    models.log_audit(user_id=g.user['id'],username=g.user['username'],action='resolve_permission_request',target_type='permission_request',target_id=request_id,detail=status,ip=request.remote_addr)
    flash('권한 요청을 처리했습니다. 권한 변경은 사용자 행에서 별도로 저장해주세요.','success')
    return redirect(url_for('admin.admin_users'))

@bp.route("/admin/users/create", methods=["POST"])
@admin_required
def admin_users_create():
    username = (request.form.get("username") or "").strip()
    display_name = (request.form.get("display_name") or "").strip()
    role = (request.form.get("role") or "staff").strip()
    password = request.form.get("password") or ""
    if not username or role not in _VALID_ROLES:
        flash("아이디와 역할을 올바르게 입력하세요.", "error")
        return redirect(url_for("admin.admin_users"))
    if len(password) < _MIN_PW_LEN:
        flash(f"비밀번호는 최소 {_MIN_PW_LEN}자 이상이어야 합니다.", "error")
        return redirect(url_for("admin.admin_users"))
    # 권한: 폼에 perm_* 가 오면 그 값, 없으면 역할 프리셋
    perms = _parse_perms_form(request.form) if any(
        k.startswith("perm_") for k in request.form) else role_preset(role)
    try:
        models.create_user(username, display_name, role, password, permissions=perms)
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("admin.admin_users"))
    models.log_audit(user_id=g.user["id"], username=g.user["username"],
                     action="create_user", target_type="user",
                     detail=f"{username} ({role})", ip=request.remote_addr)
    flash(f"'{display_name or username}' 계정을 추가했습니다.", "success")
    return redirect(url_for("admin.admin_users"))

@bp.route("/admin/users/<int:uid>/update", methods=["POST"])
@admin_required
def admin_users_update(uid):
    target = models.get_user_by_id(uid)
    if not target:
        abort(404)
    display_name = (request.form.get("display_name") or "").strip() or target["username"]
    role = (request.form.get("role") or target["role"]).strip()
    if role not in _VALID_ROLES:
        flash("올바른 역할이 아닙니다.", "error")
        return redirect(url_for("admin.admin_users"))
    perms = _parse_perms_form(request.form)
    # 사용자 관리 권한을 잃게 되는 변경이면, 다른 관리자가 최소 1명 남아야 함
    if perms.get("users", 0) < PERM_EDIT and target["perms"].get("users", 0) >= PERM_EDIT \
            and _count_user_managers(exclude_id=uid) < 1:
        flash("사용자 관리 권한을 가진 계정이 최소 1개는 있어야 합니다.", "error")
        return redirect(url_for("admin.admin_users"))
    models.update_user(uid, display_name, role, permissions=perms)
    models.log_audit(user_id=g.user["id"], username=g.user["username"],
                     action="update_user", target_type="user", target_id=uid,
                     detail=f"{target['username']} → {role} perms={perms}",
                     ip=request.remote_addr)
    flash("계정 정보·권한을 저장했습니다.", "success")
    return redirect(url_for("admin.admin_users"))

@bp.route("/admin/users/<int:uid>/password", methods=["POST"])
@admin_required
def admin_users_password(uid):
    target = models.get_user_by_id(uid)
    if not target:
        abort(404)
    password = request.form.get("password") or ""
    if len(password) < _MIN_PW_LEN:
        flash(f"비밀번호는 최소 {_MIN_PW_LEN}자 이상이어야 합니다.", "error")
        return redirect(url_for("admin.admin_users"))
    models.set_user_password(uid, password)
    models.resolve_password_reset_requests(uid, g.user["id"])
    models.log_audit(user_id=g.user["id"], username=g.user["username"],
                     action="reset_password", target_type="user", target_id=uid,
                     detail=target["username"], ip=request.remote_addr)
    flash(f"'{target['display_name'] or target['username']}' 비밀번호를 변경했습니다.", "success")
    return redirect(url_for("admin.admin_users"))

@bp.route("/admin/password-reset/<int:request_id>/resolve", methods=["POST"])
@admin_required
def admin_password_reset_resolve(request_id):
    if not models.resolve_password_reset_request(request_id, g.user["id"]):
        abort(404)
    models.log_audit(
        user_id=g.user["id"], username=g.user["username"],
        action="resolve_password_reset", target_type="password_reset_request",
        target_id=request_id, ip=request.remote_addr,
    )
    flash("비밀번호 초기화 요청을 처리 완료로 표시했습니다.", "success")
    return redirect(url_for("admin.admin_users"))

@bp.route("/admin/users/<int:uid>/active", methods=["POST"])
@admin_required
def admin_users_active(uid):
    target = models.get_user_by_id(uid)
    if not target:
        abort(404)
    activate = (request.form.get("active") == "1")
    if not activate and uid == g.user["id"]:
        flash("본인 계정은 비활성화할 수 없습니다.", "error")
        return redirect(url_for("admin.admin_users"))
    if not activate and target["perms"].get("users", 0) >= PERM_EDIT \
            and _count_user_managers(exclude_id=uid) < 1:
        flash("사용자 관리 권한을 가진 계정이 최소 1개는 있어야 합니다.", "error")
        return redirect(url_for("admin.admin_users"))
    models.set_user_active(uid, activate)
    models.log_audit(user_id=g.user["id"], username=g.user["username"],
                     action="toggle_user_active", target_type="user", target_id=uid,
                     detail=f"{target['username']} active={activate}", ip=request.remote_addr)
    flash(("활성화" if activate else "비활성화") + "했습니다.", "success")
    return redirect(url_for("admin.admin_users"))

@bp.route("/admin/users/<int:uid>/delete", methods=["POST"])
@admin_required
def admin_users_delete(uid):
    target = models.get_user_by_id(uid)
    if not target:
        abort(404)
    if uid == g.user["id"]:
        flash("본인 계정은 삭제할 수 없습니다.", "error")
        return redirect(url_for("admin.admin_users"))
    if target["perms"].get("users", 0) >= PERM_EDIT \
            and _count_user_managers(exclude_id=uid) < 1:
        flash("사용자 관리 권한을 가진 계정이 최소 1개는 있어야 합니다.", "error")
        return redirect(url_for("admin.admin_users"))
    models.delete_user(uid)
    models.log_audit(user_id=g.user["id"], username=g.user["username"],
                     action="delete_user", target_type="user", target_id=uid,
                     detail=target["username"], ip=request.remote_addr)
    flash(f"'{target['display_name'] or target['username']}' 계정을 삭제했습니다.", "success")
    return redirect(url_for("admin.admin_users"))


# ───────────────────── 엑셀 적재 (상담내역 스프레드시트 → DB) ─────────────────────
# 예전엔 NAS에 SSH로 들어가 docker exec로 tools/excel_import.py를 돌려야 했다(2026-09-16까지).
# 상담사가 매달 새 시트를 올리는 일이라 관리 메뉴에서 파일 올리기 → 시트 고르기 → 미리보기(dry-run)
# → 적재까지 되게 했다. 적재 로직은 tools/excel_import.py 그대로 쓴다(중복 판정·대기 명단 모드 포함).

def _excel_import_module():
    """tools/excel_import.py를 모듈로 불러온다. tools/에 __init__.py가 없어 경로로 읽는다."""
    import importlib.util
    path = Path(__file__).resolve().parent.parent / "tools" / "excel_import.py"
    spec = importlib.util.spec_from_file_location("bokju_excel_import", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _import_data_dir() -> Path:
    """xlsx를 두는 폴더 = DB 파일이 있는 폴더. 컨테이너에선 /data(마운트 볼륨), 개발 PC에선 저장소 루트."""
    return Path(models.DB_PATH).resolve().parent


def _safe_xlsx_name(name: str) -> str:
    """업로드 파일명 정리 — 경로 구분자·제어문자만 빼고 한글은 그대로 둔다(secure_filename은 한글을 지운다)."""
    base = os.path.basename((name or "").replace("\\", "/")).strip()
    base = "".join(ch for ch in base if ch.isprintable() and ch not in '<>:"|?*')
    if not base.lower().endswith(".xlsx"):
        base += ".xlsx"
    return base or "upload.xlsx"


def _resolve_import_file(data_dir: Path, name: str):
    """폼에서 온 파일명이 data 폴더 안의 실제 xlsx인지 확인해 Path로. 아니면 None."""
    if not name:
        return None
    candidate = (data_dir / os.path.basename(name)).resolve()
    if candidate.parent != data_dir or candidate.suffix.lower() != ".xlsx" or not candidate.is_file():
        return None
    return candidate


# 원무 입퇴원 명부 헤더 — 이 칸들이 첫 행에 있으면 상담내역이 아니라 명부로 본다.
# 상담내역 시트와 파일 모양이 완전히 달라(시트 1장, 차트번호 기준) 적재 경로도 다르다.
_ROSTER_HEADERS = ("차트번호", "수진자명", "입원일", "퇴원일")


def _roster_sheet(ws):
    """이 시트가 원무 명부면 True — 첫 행에서 필수 헤더 4개를 모두 찾는다."""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        head = {str(v).strip() for v in row if v not in (None, "")}
        return all(h in head for h in _ROSTER_HEADERS)
    return False


@bp.route("/admin/import", methods=["GET", "POST"])
@admin_required
def admin_import():
    """엑셀 적재 — 상담내역 스프레드시트(.xlsx)를 골라 시트별로 미리보기(dry-run)하고 적재한다."""
    import openpyxl
    data_dir = _import_data_dir()
    data_dir.mkdir(parents=True, exist_ok=True)
    report_text = None
    mode = None
    selected_file = request.values.get("file") or ""
    selected_sheets = request.form.getlist("sheets")

    if request.method == "POST":
        action = request.form.get("action") or ""
        if action == "upload":
            f = request.files.get("xlsx")
            if not f or not f.filename:
                flash("올릴 .xlsx 파일을 고르세요.", "error")
                return redirect(url_for("admin.admin_import"))
            fname = _safe_xlsx_name(f.filename)
            f.save(str(data_dir / fname))
            models.log_audit(user_id=g.user["id"], username=g.user["username"],
                             action="excel_import", target_type="file", target_id=None,
                             detail=f"업로드 {fname}", ip=request.remote_addr)
            flash(f"'{fname}' 을(를) 올렸습니다. 아래에서 시트를 고르고 미리보기를 누르세요.", "success")
            return redirect(url_for("admin.admin_import", file=fname))

        if action in ("roster-dryrun", "roster-apply"):
            # 원무 입퇴원 명부 — 회차 적재 후 보험유형·입원완료일·발병일 백필, 퇴원완료 전환까지 한 번에.
            path = _resolve_import_file(data_dir, selected_file)
            if not path:
                flash("파일을 찾을 수 없습니다. 목록에서 다시 고르세요.", "error")
                return redirect(url_for("admin.admin_import"))
            apply_changes = action == "roster-apply"
            close = bool(request.form.get("close_discharged"))
            create_missing = bool(request.form.get("create_missing"))
            # 하루치 스냅샷 파일이면 파일에 없는 열린 회차를 퇴원 처리(2026-09-26). 체크박스로 끌 수 있다.
            absent = bool(request.form.get("absent_discharge"))
            lines = []
            from tools.import_admission_roster import run as roster_run
            from tools.backfill_onset import read_rows as onset_rows
            try:
                summary = roster_run(str(path), apply=apply_changes, create_missing=create_missing,
                                     out=lines.append, close_discharged=close, absent_discharge=absent)
            except Exception as exc:                      # 파일 형식 오류 등
                current_app.logger.exception("명부 적재 실패")
                flash(f"명부 적재 실패: {exc}", "error")
                return redirect(url_for("admin.admin_import", file=path.name))
            # 발병일 — 명부에 있으면 회차에 채운다(멱등, 다른 값은 안 건드림)
            try:
                filled = _roster_onset(path, onset_rows, apply_changes)
                lines.append("")
                lines.append(f"발병일 — 명부에서 읽은 {filled['read']}행 중 회차 {filled['matched']}건에 기록"
                             f"{'' if apply_changes else ' (미리보기)'}")
            except Exception as exc:
                lines.append(f"발병일 처리 건너뜀: {exc}")
            report_text = "\n".join(lines)
            mode = "apply" if apply_changes else "dryrun"
            models.log_audit(
                user_id=g.user["id"], username=g.user["username"],
                action="roster_import", target_type="file", target_id=None,
                detail=f"{'적재' if apply_changes else '미리보기'} {path.name} "
                       f"행 {summary['rows']} 확정 {summary['matched']}/{summary['patients']} "
                       f"보류 {summary['held']}{' 퇴원전환' if close else ''}"
                       f"{(' 스냅샷' + summary['snapshot'] + ' 미등재퇴원 ' + str(summary['absent'])) if summary.get('snapshot') else ''}",
                ip=request.remote_addr)
            if apply_changes:
                flash(f"명부 적재 완료 — 회차 {summary['episodes']}건, 보류 {summary['held']}행. "
                      "직전 백업이 backups/pre_admission_roster_*.db 로 남았습니다.", "success")

        if action in ("dryrun", "apply"):
            path = _resolve_import_file(data_dir, selected_file)
            if not path:
                flash("파일을 찾을 수 없습니다. 목록에서 다시 고르세요.", "error")
                return redirect(url_for("admin.admin_import"))
            if not selected_sheets:
                flash("적재할 시트를 하나 이상 고르세요.", "error")
                return redirect(url_for("admin.admin_import", file=path.name))
            ei = _excel_import_module()
            wb = openpyxl.load_workbook(path, data_only=True)
            bad = [sname for sname in selected_sheets if sname not in wb.sheetnames]
            if bad:
                flash("없는 시트: " + ", ".join(bad), "error")
                return redirect(url_for("admin.admin_import", file=path.name))
            apply_changes = action == "apply"
            parts = []
            if apply_changes:
                ei.backup_db(label="excel_import")   # 시트가 여럿이어도 백업은 한 번
            totals = {"imported": 0, "updated": 0, "duplicates": 0, "skipped": 0}
            for sname in selected_sheets:
                try:
                    rep = ei.import_sheet(wb, sname, apply_changes=apply_changes, skip_backup=True)
                except SystemExit as e:      # 스키마 감지 실패 등은 SystemExit로 올라온다
                    parts.append(f"=== {sname}: 처리 불가 — {e} ===\n")
                    continue
                totals["imported"] += rep["rows_imported"]
                totals["updated"] += rep.get("rows_updated", 0)
                totals["duplicates"] += rep["rows_duplicates"]
                totals["skipped"] += rep["rows_skipped"]
                parts.append(ei.render_report(rep, apply_mode=apply_changes))
            report_text = "\n".join(parts)
            mode = "apply" if apply_changes else "dryrun"
            models.log_audit(
                user_id=g.user["id"], username=g.user["username"],
                action="excel_import", target_type="file", target_id=None,
                detail=f"{'적재' if apply_changes else '미리보기'} {path.name} [{', '.join(selected_sheets)}] "
                       f"성공 {totals['imported']} 갱신 {totals['updated']} 중복 {totals['duplicates']} 스킵 {totals['skipped']}",
                ip=request.remote_addr)
            if apply_changes:
                flash(f"적재 완료 — 성공 {totals['imported']}건, 갱신 {totals['updated']}건, "
                      f"기존 중복 {totals['duplicates']}건, 스킵 {totals['skipped']}건. 직전 백업이 backups/pre_excel_import_*.db 로 남았습니다.",
                      "success")

    files = sorted(data_dir.glob("*.xlsx"), key=lambda q: q.stat().st_mtime, reverse=True)
    is_roster = False
    file_rows = [{"name": q.name, "size_kb": round(q.stat().st_size / 1024),
                  "mtime": datetime.fromtimestamp(q.stat().st_mtime).strftime("%Y-%m-%d %H:%M")} for q in files]
    sheets = []
    path = _resolve_import_file(data_dir, selected_file)
    if path:
        ei = _excel_import_module()
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        is_roster = any(_roster_sheet(ws) for ws in wb.worksheets)
        for ws in wb.worksheets:
            # 데이터 행 수(환자이름 또는 상담일자가 있는 행) — 미리보기 전에 규모를 보여준다
            n = 0
            for row in ws.iter_rows(min_row=4, max_col=4, values_only=True):
                if any(v not in (None, "") and str(v).strip() for v in row[1:4]):
                    n += 1
            sheets.append({"name": ws.title, "rows": n, "waiting": ei.is_waiting_sheet(ws.title)})
        wb.close()
    return render_template("admin_import.html", files=file_rows, selected_file=path.name if path else "",
                           sheets=sheets, selected_sheets=selected_sheets, is_roster=is_roster,
                           report_text=report_text, mode=mode, data_dir=str(data_dir))


def _roster_onset(path, read_rows, apply_changes):
    """명부의 발병일 열 → 회차 onset_date. backfill_onset과 같은 규칙(roster_key로 찾기, 멱등)."""
    rows = read_rows(str(path))
    matched = 0
    conn = models.get_db()
    try:
        for r in rows:
            onset, chart, adm = r.get("onset_date"), r.get("chart_no"), r.get("admitted_at")
            if not onset or not chart or not adm:
                continue
            key = "%s|%s" % (chart, adm.isoformat() if hasattr(adm, "isoformat") else str(adm)[:10])
            cur = conn.execute(
                "SELECT id FROM admission_episodes WHERE roster_key = ? AND COALESCE(onset_date,'') = ''", (key,))
            hit = cur.fetchone()
            if not hit:
                continue
            matched += 1
            if apply_changes:
                conn.execute("UPDATE admission_episodes SET onset_date = ? WHERE id = ?",
                             (onset.isoformat() if hasattr(onset, "isoformat") else str(onset)[:10], hit["id"]))
        if apply_changes:
            conn.commit()
    finally:
        conn.close()
    return {"read": len(rows), "matched": matched}
