"""재원관리 → 「입원·퇴원 이력」 탭 — 기간별 입원·퇴원 명부 + 필터 + 엑셀.

근거는 models.admission_flow_events() 하나다 — 원무 명부 회차 + CRM 회차 + 상담 기록을 합쳐
(환자, 날짜, 입원/퇴원)으로 묶는다. 대시보드 '입·퇴원 현황'이 같은 함수를 쓰므로 두 화면의
명단이 어긋나지 않는다(2026-09-18 요청: 어느 화면에서도 놓치면 안 된다).
값이 엇갈리면 원무 명부가 이긴다. 상담사·나이·유입경로처럼 명부에 없는 것은 상담에서 온다.
외진(응급전원·모병원 외래치료) 복귀는 그날의 입원으로 센다(2026-09-15) — 이것도 같은 함수가
IN(source 외진) 행으로 넣어 주므로 이 모듈은 덧붙이지 않는다.

  report(args)  → {filters, rows, summary, options}
  /ward/moves.xlsx  현재 필터 그대로 엑셀
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, request, send_file

import models
from auth import login_required

bp = Blueprint("ward_moves", __name__)

KINDS = {"": "전체", "in": "입원", "out": "퇴원"}
DEFAULT_DAYS = 30


def _care_label(value) -> str:
    v = (value or "").strip()
    if not v:
        return ""
    if v.startswith("비회복"):
        return "비회복기"
    if v.startswith("회복"):
        return "회복기"
    return v


def _parse_date(s: str):
    try:
        return date.fromisoformat((s or "").strip()[:10])
    except ValueError:
        return None


def _filters(args) -> dict:
    today = date.today()
    d_to = _parse_date(args.get("date_to")) or today
    d_from = _parse_date(args.get("date_from")) or (d_to - timedelta(days=DEFAULT_DAYS - 1))
    if d_from > d_to:
        d_from, d_to = d_to, d_from
    kind = (args.get("kind") or "").strip()
    if kind not in KINDS:
        kind = ""
    return {
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(), "kind": kind,
        "doctor": (args.get("doctor") or "").strip(), "ward": (args.get("ward") or "").strip(),
        "care": (args.get("care") or "").strip(), "destination": (args.get("destination") or "").strip(),
        "reason": (args.get("reason") or "").strip(), "counselor": (args.get("counselor") or "").strip(),
        "q": (args.get("q") or "").strip(),
    }


def _load(d_from: str, d_to: str) -> list[dict]:
    """기간의 입원·퇴원 — 한 사람이 한 날 한 줄. 기간 안에 입원도 퇴원도 했으면 두 줄.

    근거는 models.admission_flow_events() 하나다(명부 회차 + CRM 회차 + 상담). 전에는 이 화면이
    명부 회차만 셌는데, 명부 적재가 늦으면 그 사이 CRM에서 입원 처리한 환자가 통째로 빠졌다
    (2026-09-18: 명부가 9/11까지만 올라와 9/14~9/18 입원 13명이 화면에 없었다).
    대시보드 '입·퇴원 현황'도 같은 함수를 쓰므로 두 화면이 어긋나지 않는다.
    """
    today = date.today()
    rows = []
    for e in models.admission_flow_events(d_from, d_to):
        adm = _parse_date(e.get("admitted_at") or "")
        dis = _parse_date(e.get("discharged_at") or "")
        age = e.get("patient_age")
        if age is None and e.get("birth_year"):
            age = today.year - int(e["birth_year"])
        rows.append({
            "episode_id": e.get("episode_id"), "consultation_id": e.get("consultation_id"),
            "patient_id": e["patient_id"], "patient_name": e.get("patient_name") or "",
            "gender": e.get("gender"), "age": age,
            "ward": e.get("ward") or "", "room": e.get("room_number") or "",
            "doctor": (e.get("attending_doctor") or "").strip(),
            "care": _care_label(e.get("care_type")),
            "dx": (e.get("diagnosis_name") or e.get("primary_diagnosis") or "").strip(),
            "admitted_at": adm.isoformat() if adm else "",
            "discharged_at": dis.isoformat() if dis else "",
            "counselor": (e.get("counselor") or "").strip(),
            "destination": (e.get("discharge_destination") or "").strip(),
            "reason": (e.get("discharge_reason") or "").strip(),
            "stay_days": ((dis or today) - adm).days + 1 if adm else None,
            "kind": e["kind"], "date": e["date"],
            "source": " · ".join(e.get("sources") or ()),
            # 외진 복귀 = 그날의 입원 — admission_flow_events가 IN(source 외진) 행으로 넣어 준다.
            "is_return": bool(e.get("is_return")), "away_type": e.get("away_type") or "",
            "away_from": e.get("away_from") or "",
        })
    rows.sort(key=lambda r: (r["date"], r["kind"] == "in", r["patient_name"]), reverse=True)
    return rows


def _match(r: dict, f: dict) -> bool:
    if f["kind"] and r["kind"] != f["kind"]:
        return False
    if f["doctor"] and r["doctor"] != f["doctor"]:
        return False
    if f["ward"] and r["ward"] != f["ward"]:
        return False
    if f["care"] and r["care"] != f["care"]:
        return False
    if f["counselor"] and r["counselor"] != f["counselor"]:
        return False
    if f["destination"] and f["destination"].casefold() not in (r["destination"] or "").casefold():
        return False
    if f["reason"] and f["reason"].casefold() not in (r["reason"] or "").casefold():
        return False
    if f["q"]:
        hay = " ".join(str(r.get(k) or "") for k in ("patient_name", "dx", "ward", "room", "destination", "reason", "doctor")).casefold()
        if f["q"].casefold() not in hay:
            return False
    return True


def _care_cells(care_in: dict, total: int) -> list[dict]:
    """수가 구분 카드용 — 회복기·비회복기는 0이어도 항상 앞에, 그 외 값·구분 없음은 뒤에."""
    order = [("회복기", "rec"), ("비회복기", "non")]
    cells = [{"label": l, "n": care_in.get(l, 0), "cls": c} for l, c in order]
    for label in sorted(care_in):
        if label not in ("회복기", "비회복기", "구분 없음"):
            cells.append({"label": label, "n": care_in[label], "cls": "etc"})
    cells.append({"label": "구분 없음", "n": care_in.get("구분 없음", 0), "cls": "none"})
    for c in cells:
        c["pct"] = round(c["n"] * 100 / total) if total else 0
    return cells


def report(args) -> dict:
    f = _filters(args)
    all_rows = _load(f["date_from"], f["date_to"])
    rows = [r for r in all_rows if _match(r, f)]
    ins = [r for r in rows if r["kind"] == "in"]
    outs = [r for r in rows if r["kind"] == "out"]
    stays = [r["stay_days"] for r in outs if r["stay_days"]]
    care_in = {}
    for r in ins:
        care_in[r["care"] or "구분 없음"] = care_in.get(r["care"] or "구분 없음", 0) + 1
    summary = {
        "admissions": len(ins), "discharges": len(outs), "net": len(ins) - len(outs),
        "avg_stay": round(sum(stays) / len(stays), 1) if stays else None,
        "long_stay": sum(1 for s in stays if s >= 365),
        "care_in": sorted(care_in.items(), key=lambda kv: -kv[1]),
        "care_cells": _care_cells(care_in, len(ins)),
        "days": (date.fromisoformat(f["date_to"]) - date.fromisoformat(f["date_from"])).days + 1,
    }
    def opts(key):
        return sorted({r[key] for r in all_rows if r.get(key)})
    options = {"doctors": opts("doctor"), "wards": opts("ward"), "cares": opts("care"), "counselors": opts("counselor")}
    active = any(f[k] for k in ("kind", "doctor", "ward", "care", "destination", "reason", "counselor", "q"))
    return {"filters": f, "rows": rows, "summary": summary, "options": options, "kinds": KINDS, "filter_active": active}


@bp.route("/ward/moves.xlsx")
@login_required
def moves_xlsx():
    """입원·퇴원 이력을 현재 필터 그대로 엑셀로 — 화면 열 순서와 같다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    rep = report(request.args)
    wb = Workbook(); ws = wb.active; ws.title = "입원·퇴원 이력"
    headers = ["연번", "날짜", "구분", "환자", "성별", "나이", "병동", "호실", "주치의", "수가", "진단", "입원일", "퇴원일", "재원일수", "퇴원 장소", "퇴원 사유", "담당 상담사"]
    ws.append(headers)
    for i, r in enumerate(rep["rows"], 1):
        ws.append([i, r["date"], "입원(복귀)" if r.get("is_return") else KINDS[r["kind"]], r["patient_name"], {"M": "남", "F": "여"}.get(r["gender"], ""),
                   r["age"] if r["age"] is not None else "", r["ward"], r["room"], r["doctor"], r["care"], r["dx"],
                   r["admitted_at"], r["discharged_at"], r["stay_days"] or "", r["destination"], r["reason"], r["counselor"]])
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E2E8F0")
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(28, len(h) * 2 + 4))
    ws.freeze_panes = "A2"
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    f = rep["filters"]
    name = f"입원퇴원이력_{f['date_from']}_{f['date_to']}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
